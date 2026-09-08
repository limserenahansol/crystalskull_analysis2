"""Example traces + FOV maps for all 18 recordings (13 Aug + 18 Aug)."""

from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

NAVY = RGBColor(0x1F, 0x3A, 0x5F)
TEAL = RGBColor(0x0D, 0x5C, 0x63)
GREY = RGBColor(0x55, 0x55, 0x55)
RED = RGBColor(0x9B, 0x1B, 0x30)

ROOT = Path(r"C:\Users\hsollim\behavior_task\crystal_skull_analysis\stim_triggered_figures")
EX = ROOT / "combined_0813_0818" / "example_neurons"
PER = EX / "per_recording"
IDX13 = ROOT / "20250813_lab_meeting" / "lab_meeting" / "recording_index.csv"
IDX18 = ROOT / "20250818_lab_meeting" / "lab_meeting" / "recording_index.csv"
OUT = EX / "CrystalSkull_face_poke_example_traces_maps_n18.pptx"


def blank(prs):
    return prs.slides.add_slide(prs.slide_layouts[6])


def set_run(p, text, size, bold, color, align=PP_ALIGN.LEFT):
    p.text = text
    p.font.size = Pt(size)
    p.font.bold = bold
    p.font.color.rgb = color
    p.font.name = "Calibri"
    p.alignment = align


def add_textbox(slide, l, t, w, h, text, size, bold, color, align=PP_ALIGN.LEFT):
    box = slide.shapes.add_textbox(Inches(l), Inches(t), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    set_run(tf.paragraphs[0], text, size, bold, color, align)
    return tf


def picture_fit(slide, img_path, left, top, max_w, max_h):
    img_path = Path(img_path)
    if not img_path.is_file():
        add_textbox(slide, left, top + max_h / 2, max_w, 0.5, "Figure missing", 16, True, RED, PP_ALIGN.CENTER)
        return
    with Image.open(img_path) as im:
        w, h = im.size
    ar = w / float(h)
    width, height = max_w, max_w / ar
    if height > max_h:
        height = max_h
        width = height * ar
    slide.shapes.add_picture(
        str(img_path),
        Inches(left + (max_w - width) / 2.0),
        Inches(top + (max_h - height) / 2.0),
        Inches(width),
        Inches(height),
    )


def load_index(path, day):
    rows = []
    lines = Path(path).read_text(encoding="utf-8").strip().splitlines()
    hdr = lines[0].split(",")
    for line in lines[1:]:
        row = dict(zip(hdr, line.split(",")))
        row["day"] = day
        rows.append(row)
    return rows


def load_examples():
    path = EX / "all_example_neurons.csv"
    if not path.is_file():
        return []
    lines = path.read_text(encoding="utf-8").strip().splitlines()
    hdr = lines[0].split(",")
    return [dict(zip(hdr, line.split(","))) for line in lines[1:]]


def title_slide(prs):
    slide = blank(prs)
    add_textbox(slide, 0.7, 1.5, 11.9, 0.9, "Example neurons from each recording", 32, True, NAVY, PP_ALIGN.CENTER)
    add_textbox(
        slide, 1.1, 2.6, 11.1, 1.1,
        "13 Aug (12) + 18 Aug (6)  |  18 movies  |  2 excited + 2 inhibited cells per FOV",
        18, False, GREY, PP_ALIGN.CENTER,
    )
    add_textbox(
        slide, 1.2, 4.0, 10.9, 2.2,
        "For each recording: strongest significant excited cell in the anterior half and posterior half, and the same for inhibited cells. Traces are EXTRACT activity with poke times. Maps show those cells on a contrast-stretched crystal-skull FOV (P left, A right).",
        18, False, TEAL, PP_ALIGN.CENTER,
    )


def two_fig_slide(prs, title, what, result, img_l, img_r):
    slide = blank(prs)
    add_textbox(slide, 0.30, 0.10, 12.7, 0.38, title, 18, True, NAVY)
    tf = add_textbox(slide, 0.30, 0.46, 12.7, 0.34, "What:  " + what, 12, False, GREY)
    p = tf.add_paragraph()
    set_run(p, "Result:  " + result, 13, True, TEAL)
    picture_fit(slide, img_l, 0.20, 0.95, 6.5, 6.3)
    picture_fit(slide, img_r, 6.65, 0.95, 6.5, 6.3)


def figure_slide(prs, title, what, result, img_path):
    slide = blank(prs)
    add_textbox(slide, 0.4, 0.12, 12.5, 0.40, title, 20, True, NAVY)
    tf = add_textbox(slide, 0.4, 0.50, 12.5, 0.36, "What:  " + what, 12, False, GREY)
    p = tf.add_paragraph()
    set_run(p, "Result:  " + result, 13, True, TEAL)
    picture_fit(slide, img_path, 0.30, 1.00, 12.7, 6.25)


def friendly(label):
    return label.replace("_", " ")


def main():
    rows = load_index(IDX13, "13 Aug") + load_index(IDX18, "18 Aug")
    examples = load_examples()
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    title_slide(prs)
    overview = EX / "example_fov_overview_n18.png"
    if overview.is_file():
        figure_slide(
            prs,
            "Where the example neurons sit   |   all 18 FOVs",
            "Each panel is one crystal-skull field after contrast stretch. Red = excited examples, blue = inhibited examples. P left, A right.",
            "The same local mix appears in every recording: example excited and inhibited cells are intermingled, with E1/I1 more anterior and E2/I2 more posterior.",
            overview,
        )

    for r in rows:
        label = r["label"]
        rec = PER / label
        pwr = r["power_mWcm2"]
        pwr_txt = "pwr NA" if pwr in ("unknown", "NA", "") else f"{pwr} mW/cm2"
        rec_ex = [e for e in examples if e.get("label") == label]
        ids = ", ".join(f"{e['tag']} #{e['neuron_id']}" for e in rec_ex) if rec_ex else "E1/E2 + I1/I2"
        two_fig_slide(
            prs,
            f"Examples  |  {r['day']}  |  {friendly(label)}  |  {r['n_pokes']} pokes",
            f"{r['mouse']}, {r['day']}, {r['session']}, {pwr_txt}, {r['mirror']}. Left: EXTRACT traces + STA. Right: those cells on the crystal-skull FOV.",
            ids,
            rec / "example_traces.png",
            rec / "example_maps.png",
        )

    prs.save(str(OUT))
    print("Wrote", OUT, "slides", len(prs.slides))


if __name__ == "__main__":
    main()
