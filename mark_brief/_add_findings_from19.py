from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Emu, Inches, Pt

NAVY = RGBColor(0x1F, 0x3A, 0x5F)
TEAL = RGBColor(0x0D, 0x5C, 0x63)
GREY = RGBColor(0x55, 0x55, 0x55)

PPT = Path(__file__).with_name("CrystalSkull_facepoke_vs_melittin_Mark_with Allen CCF.pptx")

# 1-based slide index -> (what, result). Result is the one-sentence finding.
FINDINGS = {
    19: (
        "Local % of injection cells up (red) or down (blue) vs that day's baseline, on Allen CCF.",
        "Both arms share the same cortex-wide map; morphine is a milder version (more up, less down), not a new pattern.",
    ),
    20: (
        "Local % of poke-excited (red) and poke-inhibited (blue) cells on Allen CCF. Same cells, peri-poke vs pre-poke.",
        "Poke responses are sparse: excitation is biased anterior / motor, inhibition is richer in visual cortex.",
    ),
    21: (
        "Enrichment = (area %) / (all-cortex %). Stars = FDR q<0.05 from recording-level Wilcoxon.",
        "After FDR, only MOp is significantly over-excited (1.37x); inhibition is enriched in VISl, VISam, VISp and VISrl.",
    ),
    22: (
        "% of injection cells up or down vs that day's baseline, by Allen area.",
        "Melittin up is highest in VISam / MOp / VISa / MOs; morphine raises % up in almost every area but the rank order is unchanged.",
    ),
    23: (
        "Recording-level Wilcoxon of (area rate minus that recording's overall rate), BH-FDR within each assay x direction.",
        "Face-poke excitation: MOp. Face-poke inhibition: visual areas. Melittin up: VISam, MOp, VISa, MOs, SSp-tr. Morphine: none after FDR.",
    ),
    24: (
        "Recording-level % up (melittin+morphine) minus % up (melittin). Stars would be FDR q<0.05.",
        "Morphine increases % up in almost every area (about +3 to +11 points), but 0/24 areas survive FDR (n = 4 vs 2 mice).",
    ),
}


def add_what_result(slide, what, result):
    box = slide.shapes.add_textbox(Inches(0.35), Inches(0.48), Inches(12.6), Inches(0.40))
    tf = box.text_frame
    tf.word_wrap = True
    p0 = tf.paragraphs[0]
    p0.text = "What:  " + what
    p0.font.size = Pt(12)
    p0.font.bold = False
    p0.font.color.rgb = GREY
    p0.font.name = "Calibri"
    p0.alignment = PP_ALIGN.LEFT
    p1 = tf.add_paragraph()
    p1.text = "Result:  " + result
    p1.font.size = Pt(13)
    p1.font.bold = True
    p1.font.color.rgb = TEAL
    p1.font.name = "Calibri"
    p1.alignment = PP_ALIGN.LEFT


def nudge_pictures(slide, min_top_in=0.95):
    min_top = Inches(min_top_in)
    slide_h = Inches(7.5)
    bottom_pad = Inches(0.12)
    for sh in slide.shapes:
        if sh.shape_type is None:
            continue
        # 13 = PICTURE
        if int(sh.shape_type) != 13:
            continue
        if sh.top >= min_top:
            continue
        overlap = min_top - sh.top
        new_h = sh.height - overlap
        if new_h < Inches(4.8):
            new_h = Inches(4.8)
        max_h = slide_h - min_top - bottom_pad
        if new_h > max_h:
            new_h = max_h
        sh.top = min_top
        sh.height = new_h


def main():
    prs = Presentation(str(PPT))
    n = len(prs.slides)
    for idx, (what, result) in FINDINGS.items():
        if idx > n:
            raise SystemExit(f"slide {idx} missing (deck has {n})")
        slide = prs.slides[idx - 1]
        add_what_result(slide, what, result)
        nudge_pictures(slide)
        print(f"slide {idx}: added What/Result")
    tmp = PPT.with_name(PPT.stem + "_tmp.pptx")
    prs.save(str(tmp))
    try:
        tmp.replace(PPT)
        print("saved", PPT)
    except OSError:
        print("file is open; wrote", tmp)
        print("close PowerPoint and re-run, or use the _tmp file")


if __name__ == "__main__":
    main()
