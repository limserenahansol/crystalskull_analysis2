"""One-slide PPT: face-poke excited vs inhibited cells per mouse (n=18)."""

from pathlib import Path
from shutil import copy2

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

NAVY = RGBColor(0x1F, 0x3A, 0x5F)
TEAL = RGBColor(0x0D, 0x5C, 0x63)
GREY = RGBColor(0x55, 0x55, 0x55)
RED = RGBColor(0x9B, 0x1B, 0x30)

ROOT = Path(r"C:\Users\hsollim\behavior_task\crystal_skull_analysis")
FIG = ROOT / "mark_brief" / "facepoke_per_mouse.png"
N18 = ROOT / "stim_triggered_figures" / "combined_0813_0818" / "n18"
OUTS = [
    N18 / "CrystalSkull_face_poke_n18_per_mouse.pptx",
    ROOT / "stim_triggered_figures" / "combined_0813_0818" / "CrystalSkull_face_poke_n18_per_mouse.pptx",
    ROOT / "stim_triggered_figures" / "20250813_lab_meeting_pre" / "lab_meeting" / "CrystalSkull_face_poke_n18_per_mouse.pptx",
]


def add_textbox(slide, l, t, w, h, text, size, bold, color, align=PP_ALIGN.LEFT):
    box = slide.shapes.add_textbox(Inches(l), Inches(t), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = text
    p.font.size = Pt(size)
    p.font.bold = bold
    p.font.color.rgb = color
    p.font.name = "Calibri"
    p.alignment = align
    return tf


def picture_fit(slide, img_path, left, top, max_w, max_h):
    with Image.open(img_path) as im:
        w, h = im.size
    ar = w / float(h)
    width, height = max_w, max_w / ar
    if height > max_h:
        height = max_h
        width = height * ar
    slide.shapes.add_picture(
        str(img_path),
        Inches(left + (max_w - width) / 2.0),
        Inches(top + (max_h - height) / 2.0),
        Inches(width),
        Inches(height),
    )


def main():
    if not FIG.is_file():
        raise FileNotFoundError(FIG)
    N18.mkdir(parents=True, exist_ok=True)
    copy2(FIG, N18 / "fig3_n18_per_mouse.png")

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_textbox(slide, 0.4, 0.12, 12.5, 0.38, "Face poke  |  excited vs inhibited cells per mouse   (n = 18)", 20, True, NAVY)
    tf = add_textbox(
        slide, 0.4, 0.48, 12.5, 0.32,
        "What:  Pooled within mouse. Excited/inhibited if 0.4–1.2 s after the poke differs from −2–0 s (signed-rank p<0.05).",
        12, False, GREY,
    )
    p = tf.add_paragraph()
    p.text = "Result:  Every mouse has more inhibited than excited cells: m1f 9.4% vs 22.8%, m2c 7.5% vs 30.2%, m3 12.9% vs 30.2%. Pooled 8.6% / 27.6% of 43,156 neurons; 14/18 movies p<0.05."
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = TEAL
    p.font.name = "Calibri"
    picture_fit(slide, FIG, 0.25, 0.95, 12.8, 6.3)

    blob = None
    for path in OUTS:
        path.parent.mkdir(parents=True, exist_ok=True)
        try:
            if blob is None:
                prs.save(str(path))
                blob = path.read_bytes()
            else:
                path.write_bytes(blob)
            print("Wrote", path)
        except PermissionError:
            alt = path.with_name(path.stem + "_copy.pptx")
            if blob is None:
                prs.save(str(alt))
                blob = alt.read_bytes()
            else:
                alt.write_bytes(blob)
            print("Locked, wrote", alt)


if __name__ == "__main__":
    main()
