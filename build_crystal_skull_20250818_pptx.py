"""Build the 20250818 lab-meeting PPT in the same layout as the 0813 hansol deck."""

from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

NAVY = RGBColor(0x1F, 0x3A, 0x5F)
TEAL = RGBColor(0x0D, 0x5C, 0x63)
RED = RGBColor(0x9B, 0x1B, 0x30)
GREY = RGBColor(0x55, 0x55, 0x55)

LM = Path(
    r"C:\Users\hsollim\behavior_task\crystal_skull_analysis"
    r"\stim_triggered_figures\20250818_lab_meeting\lab_meeting"
)
WORK = Path(
    r"C:\Users\hsollim\behavior_task\crystal_skull_analysis"
    r"\stim_triggered_figures\20250813_lab_meeting_pre\lab_meeting"
)
OUT_NAME = "CrystalSkull_face_poke_labmeeting_20250818.pptx"

RECS = [
    ("m1f_s00_pwrNA_mirror1", "mouse1f (m200)  s00  mirror1"),
    ("m1f_s00_pwrNA_mirror2", "mouse1f (m200)  s00  mirror2"),
    ("m2c_s00_pwrNA_mirror1", "mouse2c (m201)  s00  mirror1"),
    ("m2c_s00_pwrNA_mirror2", "mouse2c (m201)  s00  mirror2"),
    ("m2c_s01_pwrNA_mirror1", "mouse2c (m201)  s01  mirror1"),
    ("m2c_s01_pwrNA_mirror2", "mouse2c (m201)  s01  mirror2"),
]

FIGS = [
    {
        "file": "fig1_population_psth.png",
        "title": "Fig1  Population PSTH",
        "what": "Mean population activity (z-score) and fraction of active cells, aligned to each poke (t = 0 is the intensity dip).",
        "cohort_result": "Activity rises into t = 0, then drops sharply at the dip and recovers slowly — the same shape in all 6 movies.",
    },
    {
        "file": "fig2_neuron_sta.png",
        "title": "Fig2  Neuron STA",
        "what": "Each neuron's average trace around the poke, sorted by response size; right panel is the mean of excited vs inhibited cells.",
        "cohort_result": "A subset of cells go up after the dip and another subset go down; the all-cell mean is modest because they partly cancel.",
    },
    {
        "file": "fig3_responsive_cells.png",
        "title": "Fig3  Excited vs inhibited cells",
        "what": "Per neuron: mean z in 0.4–1.2 s after t=0 vs -2–0 s before; signed-rank, p<0.05, sign of the difference.",
        "cohort_result": "About 6–10% of cells are excited and 19–37% inhibited; 3 of 6 movies have a significant population response.",
    },
    {
        "file": "fig4_spatial_aligned.png",
        "title": "Fig4  Spatial poke-response map  (aligned)",
        "what": "Significant cells on each FOV after aligning anterior to the right (mirror2 flipped). Red = excited, blue = inhibited. P = posterior, A = anterior.",
        "cohort_result": "Cells are intermingled, with the same anterior-excited / posterior-inhibited bias seen on 13 Aug after this alignment.",
    },
    {
        "file": "fig5_trial_timing.png",
        "title": "Fig5  Trial raster and poke interval",
        "what": "One row per poke (population z over time) and the histogram of time between detected dips.",
        "cohort_result": "Median interval is 3.36 s. m200 has several ~20 s gaps (missed dips); m201 is nearly complete (~83–91 pokes).",
    },
]


def fmt_p(p):
    p = float(p)
    if p < 0.001:
        return f"p = {p:.1e}"
    return f"p = {p:.3f}"


def load_index(lm):
    rows = {}
    lines = (lm / "recording_index.csv").read_text(encoding="utf-8").strip().splitlines()
    hdr = lines[0].split(",")
    for line in lines[1:]:
        d = dict(zip(hdr, line.split(",")))
        rows[d["label"]] = d
    return rows


def rec_results(d, fig_file):
    n = int(d["n_neurons"])
    pk = int(d["n_pokes"])
    fe = 100 * float(d["frac_exc"])
    fi = 100 * float(d["frac_inh"])
    p = fmt_p(d["p_pop"])
    sig = "significant" if float(d["p_pop"]) < 0.05 else "not significant"
    if fig_file.endswith("psth.png"):
        return f"This FOV has {n} neurons and {pk} detected dips; the population pre vs post test is {sig} ({p})."
    if fig_file.endswith("sta.png"):
        return f"Excited {fe:.1f}% and inhibited {fi:.1f}% of {n} neurons; see the STA split on the right."
    if fig_file.endswith("cells.png"):
        return f"{fe:.1f}% excited, {fi:.1f}% inhibited out of {n} neurons; population test {sig} ({p})."
    if "spatial" in fig_file or "aligned" in fig_file:
        return f"Red = excited, blue = inhibited on this FOV ({fe:.1f}% exc, {fi:.1f}% inh of {n} neurons). P left, A right."
    if "timing" in fig_file:
        return f"{pk} dips detected in this movie; median interval ~3.36 s. Long gaps are missed dips."
    return ""


def blank(prs):
    return prs.slides.add_slide(prs.slide_layouts[6])


def set_run(p, text, size, bold, color, align=PP_ALIGN.LEFT):
    p.text = text
    p.font.size = Pt(size)
    p.font.bold = bold
    p.font.color.rgb = color
    p.font.name = "Calibri"
    p.alignment = align


def add_textbox(slide, l, t, w, h, text, size, bold, color, align=PP_ALIGN.LEFT):
    box = slide.shapes.add_textbox(Inches(l), Inches(t), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    set_run(tf.paragraphs[0], text, size, bold, color, align)
    return tf


def two_line_caption(slide, what, result):
    tf = add_textbox(slide, 0.4, 0.72, 12.5, 0.42, "What:  " + what, 14, False, GREY)
    p = tf.add_paragraph()
    set_run(p, "Result:  " + result, 15, True, TEAL)


def picture_fit(slide, img_path, left, top, max_w, max_h):
    with Image.open(img_path) as im:
        w, h = im.size
    ar = w / float(h)
    width, height = max_w, max_w / ar
    if height > max_h:
        height = max_h
        width = height * ar
    slide.shapes.add_picture(
        str(img_path),
        Inches(left + (max_w - width) / 2.0),
        Inches(top + (max_h - height) / 2.0),
        Inches(width),
        Inches(height),
    )


def method_slide(prs, title, body):
    slide = blank(prs)
    add_textbox(slide, 0.7, 2.1, 11.9, 0.7, title, 22, True, NAVY, PP_ALIGN.CENTER)
    add_textbox(slide, 1.2, 3.0, 10.9, 2.4, body, 22, False, GREY, PP_ALIGN.CENTER)


def section_slide(prs, title, subtitle):
    slide = blank(prs)
    add_textbox(slide, 0.7, 2.4, 11.9, 0.9, title, 32, True, TEAL, PP_ALIGN.CENTER)
    add_textbox(slide, 1.0, 3.5, 11.3, 1.2, subtitle, 18, False, GREY, PP_ALIGN.CENTER)


def figure_slide(prs, title, what, result, img_path):
    slide = blank(prs)
    add_textbox(slide, 0.4, 0.16, 12.5, 0.5, title, 24, True, NAVY)
    two_line_caption(slide, what, result)
    if img_path is not None and Path(img_path).is_file():
        picture_fit(slide, img_path, 0.3, 1.55, 12.7, 5.7)
    else:
        add_textbox(slide, 1, 3.5, 11, 0.6, "Figure file missing", 20, True, RED, PP_ALIGN.CENTER)


def main():
    per = LM / "per_recording"
    avg = LM / "cohort_average"
    idx = load_index(LM)

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    slide = blank(prs)
    add_textbox(slide, 0.8, 1.8, 11.7, 1.2, "Face poking on crystal skull", 36, True, NAVY, PP_ALIGN.CENTER)
    add_textbox(
        slide, 1.2, 3.2, 10.9, 1.4,
        "2025-08-18   |   t = 0 at peak_idx = intensity DIP   |   window 0.4–1.2 s vs -2–0 s",
        18, False, GREY, PP_ALIGN.CENTER,
    )
    add_textbox(
        slide, 1.2, 4.6, 10.9, 1.4,
        "Take-home: the dip-aligned movies still show a poke-locked cortex, with more inhibited than excited cells and the same anterior-excited / posterior-inhibited spatial bias after aligning mirror2.",
        18, True, TEAL, PP_ALIGN.CENTER,
    )

    method_slide(
        prs, "Experiment",
        "Whole dorsal cortex Ca2+ imaging while a motor arm pokes the face. All 6 movies on this day are poking sessions (m200 = mouse1f; m201 = mouse2c). s00/s01 are sequential recordings; power is not labeled for this day. mirror1 and mirror2 are two optical paths.",
    )
    method_slide(
        prs, "How poke time is defined",
        "No behavior camera. A blocker on the arm changes mean-frame intensity. On 13 Aug, peak_idx was a positive brightness peak. On 18 Aug, Omer marked peak_idx at the single-frame intensity DIP (the frame that goes dark). Some dips were hard to see; regular ~3.36 s spacing lets missed ones be inferred. Missed dips leave ~20 s gaps, mostly in m200.",
    )
    method_slide(
        prs, "How traces are analyzed",
        "EXTRACT traces at 33 Hz. Cut -2 to +3 s around each dip. A cell is excited or inhibited if 0.4–1.2 s differs from -2–0 s (Wilcoxon signed-rank, p<0.05). Same windows as the 13 Aug peak_idx deck.",
    )
    method_slide(
        prs, "How the spatial map is made",
        "EXTRACT spatial footprints were reconstructed from the file (ndSparse CSC). Each cell's post-pre z is painted on its footprint. Red = excited, blue = inhibited. mirror2 is flipped left-right so anterior is on the RIGHT in every panel (P left, A right — same as 13 Aug).",
    )

    section_slide(prs, "All recordings together", "6 movies  |  t = 0 is the intensity dip")
    for fig in FIGS:
        figure_slide(prs, fig["title"] + "   |   all movies", fig["what"], fig["cohort_result"], avg / fig["file"])

    for folder, rec_title in RECS:
        d = idx.get(folder)
        for fig in FIGS:
            result = rec_results(d, fig["file"]) if d else fig["cohort_result"]
            figure_slide(prs, fig["title"] + "   |   " + rec_title, fig["what"], result, per / folder / fig["file"])

    method_slide(
        prs, "Next",
        "Power for s00/s01 on this day is not labeled yet. Compare this dip-aligned deck with the 13 Aug peak_idx deck: same analysis, different t = 0 definition.",
    )

    out = LM / OUT_NAME
    prs.save(str(out))
    print("Wrote", out, "slides", len(prs.slides))
    try:
        WORK.mkdir(parents=True, exist_ok=True)
        dest = WORK / OUT_NAME
        dest.write_bytes(out.read_bytes())
        print("Copied", dest)
    except PermissionError:
        alt = WORK / "CrystalSkull_face_poke_labmeeting_20250818_copy.pptx"
        alt.write_bytes(out.read_bytes())
        print("Locked, copied", alt)


if __name__ == "__main__":
    main()
