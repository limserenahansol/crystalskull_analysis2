"""Lab-meeting PPT for old melittin BL vs injection (face-poke visual analog)."""

from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

NAVY = RGBColor(0x1F, 0x3A, 0x5F)
TEAL = RGBColor(0x0D, 0x5C, 0x63)
GREY = RGBColor(0x55, 0x55, 0x55)
RED = RGBColor(0x9B, 0x1B, 0x30)

ROOT = Path(r"C:\Users\hsollim\behavior_task\crystal_skull_analysis\melittin_lab_meeting")
AVG = ROOT / "cohort_average"
PER = ROOT / "per_recording"
EX = ROOT / "example_neurons"
EXPER = EX / "per_recording"
OUT = ROOT / "CrystalSkull_melittin_labmeeting.pptx"


def blank(prs):
    return prs.slides.add_slide(prs.slide_layouts[6])


def set_run(p, text, size, bold, color, align=PP_ALIGN.LEFT):
    p.text = text
    p.font.size = Pt(size)
    p.font.bold = bold
    p.font.color.rgb = color
    p.font.name = "Calibri"
    p.alignment = align


def add_textbox(slide, l, t, w, h, text, size, bold, color, align=PP_ALIGN.LEFT):
    box = slide.shapes.add_textbox(Inches(l), Inches(t), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    set_run(tf.paragraphs[0], text, size, bold, color, align)
    return tf


def picture_fit(slide, img_path, left, top, max_w, max_h):
    img_path = Path(img_path)
    if not img_path.is_file():
        add_textbox(slide, left, top + max_h / 2, max_w, 0.5, "Figure missing", 16, True, RED, PP_ALIGN.CENTER)
        return
    with Image.open(img_path) as im:
        w, h = im.size
    ar = w / float(h)
    width, height = max_w, max_w / ar
    if height > max_h:
        height = max_h
        width = height * ar
    slide.shapes.add_picture(
        str(img_path),
        Inches(left + (max_w - width) / 2.0),
        Inches(top + (max_h - height) / 2.0),
        Inches(width),
        Inches(height),
    )


def method_slide(prs, title, body):
    slide = blank(prs)
    add_textbox(slide, 0.7, 2.0, 11.9, 0.7, title, 22, True, NAVY, PP_ALIGN.CENTER)
    add_textbox(slide, 1.1, 2.8, 11.1, 3.4, body, 20, False, GREY, PP_ALIGN.CENTER)


def section_slide(prs, title, subtitle):
    slide = blank(prs)
    add_textbox(slide, 0.7, 2.4, 11.9, 0.9, title, 32, True, TEAL, PP_ALIGN.CENTER)
    add_textbox(slide, 1.0, 3.5, 11.3, 1.4, subtitle, 18, False, GREY, PP_ALIGN.CENTER)


def figure_slide(prs, title, what, result, img_path):
    slide = blank(prs)
    add_textbox(slide, 0.4, 0.12, 12.5, 0.40, title, 20, True, NAVY)
    tf = add_textbox(slide, 0.4, 0.50, 12.5, 0.36, "What:  " + what, 12, False, GREY)
    p = tf.add_paragraph()
    set_run(p, "Result:  " + result, 13, True, TEAL)
    picture_fit(slide, img_path, 0.30, 1.00, 12.7, 6.25)


def two_fig_slide(prs, title, what, result, img_l, img_r):
    slide = blank(prs)
    add_textbox(slide, 0.30, 0.10, 12.7, 0.38, title, 18, True, NAVY)
    tf = add_textbox(slide, 0.30, 0.46, 12.7, 0.34, "What:  " + what, 12, False, GREY)
    p = tf.add_paragraph()
    set_run(p, "Result:  " + result, 13, True, TEAL)
    picture_fit(slide, img_l, 0.20, 0.95, 6.5, 6.3)
    picture_fit(slide, img_r, 6.65, 0.95, 6.5, 6.3)


def bullets(prs, title, items):
    slide = blank(prs)
    add_textbox(slide, 0.6, 0.35, 12.1, 0.55, title, 24, True, NAVY)
    box = slide.shapes.add_textbox(Inches(0.9), Inches(1.15), Inches(11.5), Inches(5.8))
    tf = box.text_frame
    tf.word_wrap = True
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        set_run(p, "•  " + item, 18, False, GREY)
        p.space_after = Pt(14)


def load_index():
    path = ROOT / "recording_index.csv"
    if not path.is_file():
        return []
    lines = path.read_text(encoding="utf-8").strip().splitlines()
    hdr = lines[0].split(",")
    return [dict(zip(hdr, line.split(","))) for line in lines[1:]]


def load_examples():
    path = EX / "all_example_neurons.csv"
    if not path.is_file():
        return []
    lines = path.read_text(encoding="utf-8").strip().splitlines()
    hdr = lines[0].split(",")
    return [dict(zip(hdr, line.split(","))) for line in lines[1:]]


def stats(rows):
    if not rows:
        return {
            "n": 0,
            "n_sig": 0,
            "fe_mean": 0.0,
            "fi_mean": 0.0,
            "nn_sum": 0,
            "d_mean": 0.0,
        }
    p = [float(r["p_pop"]) for r in rows]
    fe = [100 * float(r["frac_up"]) for r in rows]
    fi = [100 * float(r["frac_down"]) for r in rows]
    nn = [int(r["n_neurons_inj"]) for r in rows]
    d = [float(r["d_pop_dff"]) for r in rows]
    return {
        "n": len(rows),
        "n_sig": sum(x < 0.05 for x in p),
        "fe_mean": sum(fe) / len(fe),
        "fi_mean": sum(fi) / len(fi),
        "nn_sum": sum(nn),
        "d_mean": sum(d) / len(d),
        "fe": fe,
        "fi": fi,
    }


def friendly(label):
    return label.replace("_", " ")


def main():
    rows = load_index()
    examples = load_examples()
    s = stats(rows)
    only = [r for r in rows if r.get("cohort") == "melittin_only"]
    morph = [r for r in rows if r.get("cohort") == "melittin_morphine"]
    s_only = stats(only)
    s_morph = stats(morph)

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    slide = blank(prs)
    add_textbox(slide, 0.7, 1.45, 11.9, 1.1, "Melittin on crystal skull", 36, True, NAVY, PP_ALIGN.CENTER)
    add_textbox(
        slide, 1.1, 2.75, 11.1, 1.0,
        f"Baseline vs injection  |  {s['n']} recordings  |  {s['nn_sum']:,} injection neurons",
        18, False, GREY, PP_ALIGN.CENTER,
    )
    add_textbox(
        slide, 1.2, 4.05, 10.9, 2.1,
        "Same visual language as face poking: population traces, up vs down cells, FOV maps, and example neurons. This is not poke-aligned. Baseline and injection EXTRACT cells are independent (not the same IDs).",
        18, True, TEAL, PP_ALIGN.CENTER,
    )

    method_slide(
        prs, "Experiment",
        "Whole dorsal cortex Ca2+ after melittin (bee venom) injection, with or without morphine. Each mouse has Baseline and Injection folders, mirror1path and mirror2path. CS1014 and CS0204 are melittin only; 1014m01/m02 are melittin + morphine.",
    )
    method_slide(
        prs, "Do not mix this with face poking",
        "Face poking is peri-poke on 13 Aug / 18 Aug. These movies have no poke times. Baseline and injection were extracted separately, so neuron #k in injection is not the same cell as neuron #k in baseline.",
    )
    method_slide(
        prs, "How traces and maps are made",
        "EXTRACT /T at the session frame rate (~29 Hz), binned to 1 s. Injection uses the last N seconds matched to baseline length. Maps use frr /S footprints on the injection summary image (percentile stretch + gamma + CLAHE). mirror2 is flipped so P is left and A is right.",
    )
    method_slide(
        prs, "How a cell is called up or down",
        "amp = injection-cell mean dF/F minus the mean of all baseline-cell mean dF/F. Up/down if that cell's 1 s bins differ from the baseline session mean (t-test p<0.05). Population p is a rank-sum of 1 s population mean dF/F, baseline vs injection.",
    )

    section_slide(prs, "All recordings together", "Melittin only (CS1014 + CS0204) and melittin + morphine")

    figure_slide(
        prs,
        f"Population baseline vs injection   (n = {s['n']} recordings)",
        "Each tinted line is one recording, time normalized 0-1 inside the matched window. Blue = baseline mean, orange = injection mean.",
        f"{s['n_sig']}/{s['n']} recordings have population p<0.05. Mean session Delta dF/F = {s['d_mean']:+.4f}.",
        AVG / "fig1_population_psth.png",
    )
    figure_slide(
        prs,
        "Injection up vs down cells   (mean of recordings)",
        "Left analog of the poke STA: mean z of injection cells classified up or down vs the baseline session mean. Black = all injection cells.",
        "The two groups have opposite sign by construction. The all-cell mean is the residual after they cancel.",
        AVG / "fig2_neuron_sta.png",
    )
    figure_slide(
        prs,
        "Up vs down fractions   (n recordings)",
        "An injection cell is up or down if its 1 s dF/F bins differ from the baseline session mean (p<0.05). One dot = one recording.",
        f"Movie-mean {s['fe_mean']:.1f}% up and {s['fi_mean']:.1f}% down. {s['n_sig']}/{s['n']} recordings p<0.05 at the population level.",
        AVG / "fig3_responsive_cells.png",
    )
    figure_slide(
        prs,
        "Injection FOV maps  (cannot pixel-average)",
        "EXTRACT footprints x amp. Red = up vs baseline mean, blue = down. P left, A right after flipping mirror2. FOVs are not CCF aligned.",
        "The same local mix appears across recordings. There is no single average cortex image.",
        AVG / "fig4_spatial_response.png",
    )
    figure_slide(
        prs,
        "Cell counts and session Delta",
        "Left: EXTRACT neuron count in each injection movie. Right: population mean dF/F (injection minus baseline).",
        "Independent cell counts. Use the Delta bars, not matched-cell before/after.",
        AVG / "fig5_session_stats.png",
    )
    figure_slide(
        prs,
        "Take-home  |  session Delta dF/F",
        "One bar per recording. Teal = melittin only. Orange = melittin + morphine.",
        f"Melittin only n={s_only['n']}, mean Delta {s_only['d_mean']:+.4f}. Melittin+morphine n={s_morph['n']}, mean Delta {s_morph['d_mean']:+.4f}.",
        AVG / "fig_takehome_delta.png",
    )
    figure_slide(
        prs,
        "Take-home  |  paired CS1014 mice",
        "Same biological mouse and mirror: CS1014-1b with 1014m01b, CS1014-2w with 1014m02w. CS0204 has no morphine pair in this pack.",
        "Lines connect the two drug conditions. Small n; treat as exploratory.",
        AVG / "fig_takehome_paired.png",
    )

    overview = EX / "example_fov_overview.png"
    if overview.is_file():
        figure_slide(
            prs,
            "Where the example neurons sit   |   all injection FOVs",
            "Each panel is one contrast-stretched crystal-skull field. Red = up examples, blue = down examples. P left, A right.",
            "Example up and down cells are intermingled. U1/D1 are more anterior; U2/D2 more posterior.",
            overview,
        )

    section_slide(prs, "Example neurons", "Two up + two down injection cells per recording (not paired to baseline)")

    for r in rows:
        label = r["label"]
        rec = EXPER / label
        rec_ex = [e for e in examples if e.get("label") == label]
        ids = ", ".join(f"{e['tag']} #{e['neuron_id']}" for e in rec_ex) if rec_ex else "U1/U2 + D1/D2"
        two_fig_slide(
            prs,
            f"Examples  |  {friendly(label)}  |  {r.get('cohort', '')}",
            f"{r['mouse']}, {r['mirror']}. Left: injection EXTRACT traces. Right: those cells on the injection FOV. Not the same IDs as baseline.",
            ids,
            rec / "example_traces.png",
            rec / "example_maps.png",
        )

    bullets(
        prs,
        "Interpretation",
        [
            "This pack answers the same visual questions as face poking, but the clock is the recording, not a poke.",
            "Up/down is each injection cell versus the baseline session mean, not the same cell before versus after.",
            "Population Delta dF/F is the honest mouse/plane comparison (injection minus baseline).",
            "Morphine pairing is only CS1014-1b/2w vs 1014m01b/02w. CS0204 is extra melittin-only.",
        ],
    )
    bullets(
        prs,
        "What we should not over-claim",
        [
            "FOVs are not CCF/bregma registered. Anterior means the right side after the mirror2 flip.",
            "Do not report a matched-cell trajectory from baseline to injection; EXTRACT IDs were not matched.",
            "dF/F here uses an 8th-percentile baseline on 1 s bins of EXTRACT /T, not the slow rolling F0 of extract_analysis2.",
            "Do not mix these numbers with the 13 Aug / 18 Aug face-poke PSTHs.",
        ],
    )

    ROOT.mkdir(parents=True, exist_ok=True)
    try:
        prs.save(str(OUT))
        print("Wrote", OUT, "slides", len(prs.slides))
    except PermissionError:
        alt = ROOT / "CrystalSkull_melittin_labmeeting_copy.pptx"
        prs.save(str(alt))
        print("Locked original, wrote", alt, "slides", len(prs.slides))


if __name__ == "__main__":
    main()
