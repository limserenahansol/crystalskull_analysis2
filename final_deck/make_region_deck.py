"""The two new region slides, plus a methods supplement for them.

Two outputs, deliberately:
  CrystalSkull_ADD_region_slides.pptx       the 3 new slides on their own, to drop
                                            into the existing deck wherever wanted
  CrystalSkull_summary_7slides_for_Mark.pptx  the whole thing in one file

The original CrystalSkull_summary_5slides_for_Mark.pptx is left exactly as it is.

python-pptx only. Never PowerPoint COM: it closes the documents Hansol has open.
"""
from __future__ import annotations

import json
import os

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Emu, Inches, Pt

HERE = os.path.dirname(os.path.abspath(__file__))
FIG = os.path.join(HERE, "figures")
ST = json.load(open(os.path.join(HERE, "region_stats.json")))

W, H = Inches(13.333), Inches(7.5)
NAVY = RGBColor(0x1F, 0x3A, 0x5F)
INK = RGBColor(0x22, 0x22, 0x22)


def new_deck():
    prs = Presentation()
    prs.slide_width, prs.slide_height = W, H
    return prs


def tb(slide, left, top, width, height, text, size=13, bold=False, color=INK,
       align=PP_ALIGN.LEFT, space_after=4):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    for i, line in enumerate(text.split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_after = Pt(space_after)
        run = p.add_run()
        run.text = line
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.color.rgb = color
        run.font.name = "Calibri"
    return box


def figure_slide(prs, title, png, caption, notes):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    tb(s, Inches(.45), Inches(.22), Inches(12.45), Inches(.6), title,
       size=24, bold=True, color=NAVY)
    ph = Inches(4.72)
    pw = Emu(int(ph * (13.0 / 5.5)))
    s.shapes.add_picture(os.path.join(FIG, png), Emu(int((W - pw) / 2)),
                         Inches(.82), width=pw, height=ph)
    tb(s, Inches(.55), Inches(5.66), Inches(12.25), Inches(1.72), caption,
       size=12.5, color=INK, space_after=3)
    s.notes_slide.notes_text_frame.text = notes
    return s


def text_slide(prs, title, blocks, notes=""):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    tb(s, Inches(.45), Inches(.22), Inches(12.45), Inches(.6), title,
       size=22, bold=True, color=NAVY)
    for k, body in enumerate(blocks):
        tb(s, Inches(.52 + k * 6.35), Inches(1.00), Inches(6.1), Inches(6.1), body,
           size=11.5, color=INK, space_after=3)
    if notes:
        s.notes_slide.notes_text_frame.text = notes
    return s


A1, A2 = "CS1014-1b", "CS1014-2w"
C, D = ST["C"], ST["D"]

# ------------------------------------------------------------------ content
MEL_TITLE = "3.  Melittin: posterior-medial parietal and motor cortex"
MEL_PNG = "S6_melittin_regions.png"
MEL_CAP = (
    "What we did:  ran exactly the face-poke region analysis on the melittin-only arm "
    "(81,834 cells, 4 mice), expressing each Allen area's response rate as a fold "
    "enrichment over that arm's own cortex-wide rate so the two paradigms can be read "
    "side by side.\n"
    "What we found:  melittin drives posterior-medial parietal cortex (VISam 1.74×, "
    "VISpm 1.46×, VISa 1.42×) and motor cortex (MOp 1.42×, MOs 1.36×), and spares "
    "lateral temporal and auditory areas (TEa, VISpor, AUDp all ≈ 0.22–0.24×). 'Up' "
    "and 'down' are near-perfect mirrors of each other (ρ = −0.98), so they carry one "
    "map and not two; and face poke's posterior inhibition gradient is not "
    "demonstrable here."
)
MEL_NOTES = (
    "WHY THIS SLIDE MIRRORS SLIDE 2 EXACTLY\n"
    "Same enrichment recipe, same >= 300 cell floor, same binomial + BH-FDR, same "
    "colour meanings. The only thing that differs is which arm the cells come from. "
    "That is deliberate: it is the only way the two paradigms can be compared by eye.\n\n"
    "PANEL C IS THE IMPORTANT ONE\n"
    "The up-enrichment and down-enrichment maps correlate at Spearman rho = -0.98 "
    "(p = 1e-19, 27 areas). They are the same map with the sign flipped, so reporting "
    "both is reporting one number twice.\n\n"
    "The reason is in the base rates: 28.3% of cells are called up and 63.0% down, "
    "which is 91.3% of all cells called responsive. The melittin test uses time bins as "
    "its sample, so df is in the thousands and nearly every cell crosses p<0.05 in one "
    "direction or the other. An area that is high on 'up' is therefore mechanically low "
    "on 'down'. Supplement S2 has the full argument.\n\n"
    "This is why the absolute 28% / 63% must never be compared with the face-poke "
    "8.6% / 27.6% as biology. The enrichment ratios ARE comparable, because each is "
    "divided by its own arm's cortex-wide rate - that is what slide 6 (dissociation) "
    "rests on.\n\n"
    "PANEL D - NO GRADIENT, AND WHY THAT IS NOT THE SAME AS 'NO EFFECT'\n"
    "Face poke: +7.2 pp posterior, 18/18 recordings, p = 7.6e-06.\n"
    "Melittin:  +2.2 pp posterior, 3/4 mice, p = 0.38.\n"
    "Two separate reasons not to read the melittin result as a demonstrated absence:\n"
    "  1. the effect is 3x smaller, and\n"
    "  2. a two-sided Wilcoxon signed-rank on n = 4 has a FLOOR of p = 0.125. Even if "
    "all 4 animals had agreed, this test could not have reached p < 0.05. It is "
    "structurally incapable of significance at this n.\n"
    "Running it over the 8 recordings instead of the 4 mice would clear that floor, but "
    "mirror1 and mirror2 are two recordings of the SAME session covering the same "
    "bilateral FOV - they are technical replicates, and using them as n = 8 is "
    "pseudoreplication. The honest statement is 'not demonstrable', not 'absent'.\n\n"
    "PER-MOUSE CHECK\n"
    "CS0204-1b 29.7% up, CS0204-2w 20.9%, CS1014-1b 28.5%, CS1014-2w 28.9%. CS0204-2w "
    "sits low and also has the fewest cells (7,256), so it carries the least weight in "
    "the pooled number - the pooled 28.3% is not driven by one animal."
)

MOR_TITLE = "5.  Is morphine region-specific?  Not reproducibly, at n = 2"
MOR_PNG = "S7_morphine_regions.png"
MOR_CAP = (
    "What we did:  asked the question in the two layers it needs — first, within each "
    "of the two animals recorded under both arms, does the morphine effect differ "
    "between Allen areas (Cochran's Q on the per-area log odds ratio); second, do the "
    "two animals agree on which areas.\n"
    f"What we found:  within each animal the effect is strongly region-dependent "
    f"(I² = {C[A1]['I2']:.0f} % and {C[A2]['I2']:.0f} %, p = 1e-16 and 1e-46), but the "
    f"two animals rank the areas completely differently "
    f"(r = {D['pearson_lor'][0]:+.2f}, p = {D['pearson_lor'][1]:.2f}; "
    f"{D['agree']}/{D['n']} directions agree, which is chance). Only the whole-cortex "
    f"shift reproduces: OR 1.15 and 1.16, both above 1. No individual area can be "
    f"named from n = 2."
)
MOR_NOTES = (
    "FIRST, THE DESIGN CONSTRAINT\n"
    "There is NO morphine-alone arm and NO morphine + face-poke arm in this dataset. "
    "Morphine only ever appears as melittin + morphine, on 20260408, in two animals: "
    "1014m01b and 1014m02w, which Hansol confirmed are CS1014-1b and CS1014-2w from the "
    "melittin-only cohort. So every morphine statement is a within-animal melittin -> "
    "melittin+morphine comparison in those two mice.\n\n"
    "WHY TWO LAYERS AND NOT ONE\n"
    "Panel A alone would support a headline like 'morphine acts region-specifically' - "
    "and the statistics for it are overwhelming, Q = 129 and 278. But the sample in "
    "that test is CELLS within one animal. It is the right test for 'do regions differ "
    "inside this animal', and it says nothing at all about whether the pattern is a "
    "property of morphine or a property of that animal's particular FOV and cell "
    "sampling. Panel B is what distinguishes those, and it comes out flat.\n\n"
    "THE NUMBERS\n"
    f"{A1}: pooled log OR {C[A1]['pool']:+.3f} (OR "
    f"{2.718281828**C[A1]['pool']:.3f}), Q = {C[A1]['Q']:.0f}, df = {C[A1]['df']}, "
    f"p = {C[A1]['pQ']:.1e}, I2 = {C[A1]['I2']:.0f}%, "
    f"{C[A1]['n_sig']}/{C[A1]['df']+1} areas individually significant after BH.\n"
    f"{A2}: pooled log OR {C[A2]['pool']:+.3f} (OR "
    f"{2.718281828**C[A2]['pool']:.3f}), Q = {C[A2]['Q']:.0f}, df = {C[A2]['df']}, "
    f"p = {C[A2]['pQ']:.1e}, I2 = {C[A2]['I2']:.0f}%, "
    f"{C[A2]['n_sig']}/{C[A2]['df']+1} areas individually significant after BH.\n"
    f"Cross-animal: Pearson r = {D['pearson_lor'][0]:+.2f} (p = "
    f"{D['pearson_lor'][1]:.2f}), Spearman rho = {D['spearman_lor'][0]:+.2f} "
    f"(p = {D['spearman_lor'][1]:.2f}), {D['agree']}/{D['n']} directions agree "
    f"(binomial vs 0.5, p = {D['p_binom']:.2f}), CIs overlap in "
    f"{D['ci_overlap']}/{D['n']}.\n\n"
    "WHAT THE FOREST PLOT SHOWS IF SOMEONE LOOKS CLOSELY\n"
    "The areas at the top (AUDpo, RSPd, VISal, RSPagl, VISl) and at the bottom (SSp-n, "
    "SSp-ll, SSp-un) are ordered by the MEAN of the two animals, which is why they look "
    "tidy. Look at the individual bars and the two colours disagree constantly: AUDd is "
    "+0.75 in one animal and -0.09 in the other, SSp-ll is +0.20 and -0.64, SSp-n is "
    "-0.50 and +0.04. That disagreement is panel B.\n\n"
    "VISpor in CS1014-2w has a CI running past +/- 4 because its up-rate is near zero; "
    "the axis is clipped at +/- 1.6 and three CI ends are marked with arrowheads.\n\n"
    "A THIRD ANALYSIS, ALSO NEGATIVE\n"
    "Does morphine act most where melittin already acts most? Correlating each area's "
    "melittin-alone up-rate against its morphine-minus-melittin change gives r = -0.51 "
    "in CS1014-1b and -0.23 in CS1014-2w. A negative correlation is EXPECTED here even "
    "under no real effect, because the noise in x re-enters y with a minus sign. I "
    "simulated that null (2,000 draws, both arms resampled at the pooled per-area rate): "
    "the null mean is about -0.10 with a 95% range of roughly -0.48 to +0.31. "
    "CS1014-1b's -0.51 sits at the 2nd percentile, outside the artifact; CS1014-2w's "
    "-0.23 sits at the 27th, inside it. One animal yes, one no - the same n = 2 problem, "
    "so it is not on the slide as a finding.\n\n"
    "WHAT WOULD SETTLE IT\n"
    "Two more paired animals. With n = 4, the cross-animal correlation of per-area log "
    "ORs becomes interpretable, and a sign test on the whole-cortex shift can reach "
    "p = 0.125; with n = 6 it can reach p = 0.031."
)

SUPP_TITLE = "S6.  Methods — the region-level tests on slides 3 and 5"
SUPP_BLOCKS = [
    "ENRICHMENT  (slides 2, 3, 6)\n"
    "Fold enrichment = that Allen area's response rate ÷ the cortex-wide rate of the "
    "SAME arm. Areas with ≥ 300 cells only. Two-sided binomial test of each area "
    "against the cortex-wide rate, Benjamini-Hochberg FDR across the areas tested "
    "(19 for face poke, 27 for melittin).\n"
    "Dividing by each arm's own base rate is what makes the two paradigms comparable "
    "despite base rates of 8.6 % and 28.3 % — see supplement S2.\n"
    "\n"
    "A–P GRADIENT  (slides 2, 3)\n"
    "Each recording is split at its OWN median AP_mm, so the result does not depend on "
    "how far forward that FOV happened to sit. Wilcoxon signed-rank on the paired "
    "anterior/posterior fractions.\n"
    "Unit of analysis: face poke n = 18 recordings; melittin n = 4 mice, because "
    "mirror1 and mirror2 are two recordings of the SAME session covering the same "
    "bilateral FOV. They are technical replicates, not independent samples.\n"
    "At n = 4 a two-sided signed-rank test floors at p = 0.125, so it cannot reach "
    "significance however consistent the animals are. 'Not demonstrable' ≠ 'absent'.\n"
    "\n"
    "WITHIN-ANIMAL REGION HETEROGENEITY  (slide 5, panel A)\n"
    "Per isocortical area with ≥ 100 cells in both arms, the 2×2 table (up / not-up) × "
    "(melittin / melittin+morphine) gives a log odds ratio and its SE, with a "
    "Haldane–Anscombe 0.5 correction so a zero count cannot break the log.\n"
    "Inverse-variance pooling gives that animal's overall log OR.\n"
    "Cochran's Q = Σ wᵢ(logORᵢ − logOR_pool)², wᵢ = 1/SEᵢ², df = K−1, χ². I² = "
    "(Q−df)/Q.\n"
    "A significant Q means the effect differs between areas.",

    "WHAT COCHRAN'S Q DOES AND DOES NOT LICENSE\n"
    "The sample in that test is CELLS inside one animal. That is the right unit for "
    "asking whether regions differ within that animal, and the wrong unit for asking "
    "whether morphine is region-specific in general. Treating cells as independent "
    "animals would turn Q = 129 into a false headline — which is why slide 5 has a "
    "panel B.\n"
    "\n"
    "CROSS-ANIMAL REPRODUCIBILITY  (slide 5, panel B)\n"
    "Per-area log ORs from the two animals, same 22 areas, correlated with Pearson and "
    "Spearman. Direction agreement tested against chance with a two-sided binomial "
    "(null 0.5). 95 % CI overlap counted per area.\n"
    "All four read the same way: r = +0.10, ρ = +0.10, 11/22 directions, CIs overlap "
    "in 14/22.\n"
    "\n"
    "THE REGRESSION-TO-THE-MEAN NULL  (not on the slide)\n"
    "Correlating an area's melittin rate (x) against its morphine-minus-melittin "
    "change (y) is biased negative by construction: the noise in x re-enters y with a "
    "minus sign. Both arms were resampled 2,000× at each area's pooled rate "
    "(default_rng(0)) to size that bias. Null mean ≈ −0.10, 95 % range ≈ −0.48 to "
    "+0.31; the observed values sit at the 2nd and 27th percentile.\n"
    "\n"
    "DESIGN LIMIT, STATED ONCE\n"
    "No morphine-alone arm, no morphine + face-poke arm. Morphine exists only as "
    "melittin + morphine, in two animals. Every morphine number inherits that n = 2.",
]
SUPP_NOTES = (
    "This slide exists so slides 3 and 5 can stay readable. If Mark asks 'how did you "
    "decide an area is enriched' or 'why is n = 4 and not 8', the answer is here."
)

# ------------------------------------------------------------- the add-on deck
add = new_deck()
figure_slide(add, MEL_TITLE, MEL_PNG, MEL_CAP, MEL_NOTES)
figure_slide(add, MOR_TITLE, MOR_PNG, MOR_CAP, MOR_NOTES)
text_slide(add, SUPP_TITLE, SUPP_BLOCKS, SUPP_NOTES)
p_add = os.path.join(HERE, "CrystalSkull_ADD_region_slides.pptx")
add.save(p_add)
print(f"wrote {p_add}  ({len(add.slides._sldIdLst)} slides)")

# ------------------------------------------------------------ the merged deck
# Captions and notes are READ back out of the existing 5-slide file rather than
# re-derived, so the two decks cannot drift. Do not import make_final_deck here:
# building it is that module's import side effect, and it fails outright when the
# file is open in PowerPoint.

MAIN = [
    ("1.  All 30 recordings registered to the Allen CCF — no manual landmarks",
     "S1_registration.png", 0),
    ("2.  Face poke: somatosensory excitation, posterior inhibition",
     "S2_facepoke.png", 1),
    (MEL_TITLE, MEL_PNG, None),
    ("4.  Morphine does not silence the melittin response (within animal)",
     "S3_morphine.png", 2),
    (MOR_TITLE, MOR_PNG, None),
    ("6.  Face poke and melittin recruit different cortical territory",
     "S4_dissociation.png", 3),
    ("7.  Face video is indexed and synced; grimace scoring needs a rig change",
     "S5_behavior.png", 4),
]

src = Presentation(os.path.join(HERE, "CrystalSkull_summary_5slides_for_Mark.pptx"))
src_slides = list(src.slides)


def caption_of(i):
    s = src_slides[i]
    boxes = [sh for sh in s.shapes if sh.has_text_frame]
    return boxes[-1].text_frame.text


def notes_of(i):
    return src_slides[i].notes_slide.notes_text_frame.text


merged = new_deck()
for title, png, srci in MAIN:
    if srci is None:
        cap = MEL_CAP if png == MEL_PNG else MOR_CAP
        nts = MEL_NOTES if png == MEL_PNG else MOR_NOTES
    else:
        cap, nts = caption_of(srci), notes_of(srci)
    figure_slide(merged, title, png, cap, nts)

for i in range(5, len(src_slides)):
    s = src_slides[i]
    boxes = [sh for sh in s.shapes if sh.has_text_frame]
    text_slide(merged, boxes[0].text_frame.text,
               [b.text_frame.text for b in boxes[1:]],
               s.notes_slide.notes_text_frame.text if s.has_notes_slide else "")
text_slide(merged, SUPP_TITLE, SUPP_BLOCKS, SUPP_NOTES)

p_all = os.path.join(HERE, "CrystalSkull_summary_7slides_for_Mark.pptx")
merged.save(p_all)
print(f"wrote {p_all}  ({len(merged.slides._sldIdLst)} slides)")
