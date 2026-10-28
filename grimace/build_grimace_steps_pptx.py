"""Slides that show the three GrimACE steps on our frames, and which step fails."""
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Emu, Inches, Pt

FIG = Path(__file__).resolve().parent / "grimace_output" / "figures"
OUT = Path(__file__).resolve().parent / "GrimACE_three_steps_check_EN.pptx"

NAVY = RGBColor(0x1F, 0x3A, 0x5F)
TEAL = RGBColor(0x0D, 0x5C, 0x63)
RED = RGBColor(0x9B, 0x1B, 0x30)
BLACK = RGBColor(0x22, 0x22, 0x22)

SLIDES = [
    {
        "title": "1. Paper reference: a working GrimACE",
        "image": "paper_fig1_grimace_system.png",
        "summary": "The mouse stands in a small infrared box. The face is a small part of the frame. The app then crops that face and prints five scores: orbital, nose, cheek, ear, whisker.",
        "conclusion": "This is the picture the quality network and the YOLO were trained on. Our camera is already a tight close-up, so it is a different image.",
        "color": TEAL,
        "credit": "Sturman et al., Lab Animal 2026, Fig. 1. CC BY 4.0. doi.org/10.1038/s41684-026-01695-9",
    },
    {
        "title": "2. Paper reference: a real pain increase",
        "image": "paper_fig2_good_assay.png",
        "summary": "After craniotomy with a head implant, expert and automated MGS rise from about 0.3 before surgery to about 1.1 at 4 hours, then fall. Panel a is the front face the model expects.",
        "conclusion": "A successful GrimACE result in this paper is a within-animal rise of about 1 point. Scores stay near 1. They do not jump to 5-8.",
        "color": TEAL,
        "credit": "Sturman et al., Lab Animal 2026, Fig. 2. CC BY 4.0. doi.org/10.1038/s41684-026-01695-9",
    },
    {
        "title": "3. The same three steps on one of our frames",
        "image": "figS1_three_steps_one_frame.png",
        "summary": "Quality score, then the YOLO box, then the vision transformer on that crop. The orange box is the model coordinate, not a hand drawing.",
        "conclusion": "YOLO confidence is 0.45, but the box is on the paws. The transformer scores that paw crop and returns MGS 2.",
        "color": RED,
        "credit": "",
    },
    {
        "title": "4. Quality network on our camera",
        "image": "figS2_quality_network.png",
        "summary": "Six sharp front faces from our recordings score 0.21 to 0.37. The network scale is 0 to 1.",
        "conclusion": "The quality network does not separate good faces from bad ones on this camera. It would reject almost every frame.",
        "color": RED,
        "credit": "",
    },
    {
        "title": "5. YOLO face box on our camera",
        "image": "figS3_yolo_boxes.png",
        "summary": "Orange boxes passed the 0.10 cutoff and were used for scoring. Red labels are frames with no box, so the center of the image was cropped instead.",
        "conclusion": "When a box appears, it sits on the paws or the top of the head. Clear front-face frames get no box.",
        "color": RED,
        "credit": "",
    },
    {
        "title": "6. Vision transformer: a score for the whole crop",
        "image": "figS4_vit_scores.png",
        "summary": "The transformer does not mark the eye, nose, cheek, ear, or whiskers. It returns 0, 1, 2, or not rateable for the entire crop.",
        "conclusion": "A crop that contains only the paws still receives an ear score of 1. A number is not evidence that the feature was found.",
        "color": RED,
        "credit": "",
    },
    {
        "title": "7. Which step works on crystal-skull video",
        "image": "figS5_which_step_fails.png",
        "summary": "Against the paper's front-face box image, our quality scores stay low, the best YOLO box is on the paws, and the transformer still emits an MGS.",
        "conclusion": "None of the three steps is reliable on these recordings. This is not yet a pain assay for the crystal skull.",
        "color": RED,
        "credit": "",
    },
]


def _add_text(slide, left, top, width, height, text, size, color, label="", bold=False):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.LEFT
    if label:
        lab = p.add_run()
        lab.text = label
        lab.font.size = Pt(size)
        lab.font.bold = True
        lab.font.color.rgb = color
        lab.font.name = "Calibri"
    run = p.add_run()
    run.text = text
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = "Calibri"


def _fit(path: Path, max_w: int, max_h: int) -> tuple[int, int]:
    with Image.open(path) as im:
        w, h = im.size
    scale = min(max_w / w, max_h / h)
    return int(w * scale), int(h * scale)


def main() -> None:
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]
    max_w, max_h = Inches(12.5), Inches(3.85)
    for spec in SLIDES:
        slide = prs.slides.add_slide(blank)
        _add_text(slide, Inches(0.38), Inches(0.12), Inches(12.6), Inches(0.48),
                  spec["title"], 32, NAVY, bold=True)
        img = FIG / spec["image"]
        if not img.is_file():
            raise FileNotFoundError(img)
        nw, nh = _fit(img, max_w, max_h)
        left = int((prs.slide_width - nw) / 2)
        slide.shapes.add_picture(str(img), Emu(left), Inches(0.62), Emu(nw), Emu(nh))
        _add_text(slide, Inches(0.4), Inches(4.62), Inches(12.5), Inches(1.15),
                  spec["summary"], 20, BLACK, label="Summary.  ")
        _add_text(slide, Inches(0.4), Inches(5.85), Inches(12.5), Inches(1.15),
                  spec["conclusion"], 22, spec["color"], label="Conclusion.  ")
        if spec.get("credit"):
            _add_text(slide, Inches(0.4), Inches(7.08), Inches(12.5), Inches(0.32),
                      spec["credit"], 12, RGBColor(0x66, 0x66, 0x66))
    prs.save(OUT)
    print("wrote", OUT)


if __name__ == "__main__":
    main()
