"""One-slide PPT: paired CS1014 FOV maps, melittin vs morphine only."""

from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

NAVY = RGBColor(0x1F, 0x3A, 0x5F)
TEAL = RGBColor(0x0D, 0x5C, 0x63)
GREY = RGBColor(0x55, 0x55, 0x55)

ROOT = Path(r"C:\Users\hsollim\behavior_task\crystal_skull_analysis")
FIG = ROOT / "mark_brief" / "fig4_paired_CS1014_melittin_vs_morphine.png"
OUTS = [
    ROOT / "mark_brief" / "CrystalSkull_melittin_vs_morphine_paired_CS1014.pptx",
    ROOT / "melittin_lab_meeting" / "CrystalSkull_melittin_vs_morphine_paired_CS1014.pptx",
]


def add_textbox(slide, l, t, w, h, text, size, bold, color):
    box = slide.shapes.add_textbox(Inches(l), Inches(t), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = text
    p.font.size = Pt(size)
    p.font.bold = bold
    p.font.color.rgb = color
    p.font.name = "Calibri"
    p.alignment = PP_ALIGN.LEFT
    return tf


def picture_fit(slide, img_path, left, top, max_w, max_h):
    with Image.open(img_path) as im:
        w, h = im.size
    ar = w / float(h)
    width, height = max_w, max_w / ar
    if height > max_h:
        height = max_h
        width = height * ar
    slide.shapes.add_picture(
        str(img_path),
        Inches(left + (max_w - width) / 2.0),
        Inches(top + (max_h - height) / 2.0),
        Inches(width),
        Inches(height),
    )


def main():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_textbox(
        slide, 0.35, 0.08, 12.6, 0.36,
        "Melittin vs melittin + morphine  |  the two mice that have both  (CS1014)",
        18, True, NAVY,
    )
    tf = add_textbox(
        slide, 0.35, 0.42, 12.6, 0.30,
        "What:  Same biological mouse and mirror. Left of each pair = melittin only. Right = melittin + morphine. Red = up vs BL mean, blue = down. EXTRACT IDs are not matched across days.",
        12, False, GREY,
    )
    p = tf.add_paragraph()
    p.text = "Result:  3/4 pairs shift toward less suppression or net increase with morphine (CS1014-1b mirror2 is the exception). CS0204 is omitted; those mice have no morphine day."
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = TEAL
    p.font.name = "Calibri"
    picture_fit(slide, FIG, 0.15, 0.88, 13.0, 6.45)

    blob = None
    for path in OUTS:
        path.parent.mkdir(parents=True, exist_ok=True)
        try:
            if blob is None:
                prs.save(str(path))
                blob = path.read_bytes()
            else:
                path.write_bytes(blob)
            print("Wrote", path)
        except PermissionError:
            alt = path.with_name(path.stem + "_copy.pptx")
            prs.save(str(alt))
            print("Locked, wrote", alt)


if __name__ == "__main__":
    main()
