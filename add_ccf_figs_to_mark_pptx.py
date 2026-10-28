"""Append CCF/region figures to the existing Mark brief. Does not rebuild or replace slides."""

from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

NAVY = RGBColor(0x1F, 0x3A, 0x5F)
GREY = RGBColor(0x55, 0x55, 0x55)

ROOT = Path(r"C:\Users\hsollim\behavior_task\crystal_skull_analysis")
SRC = ROOT / "mark_brief" / "CrystalSkull_facepoke_vs_melittin_Mark.pptx"
FIG = ROOT / "ccf_registration" / "figures_v2"

FIGURES = [
    ("Allen CCF  |  all recordings merged", FIG / "fig_merged_ccf_maps.png"),
    ("Face poke  |  region enrichment vs rest of cortex", FIG / "fig_poke_region_enrichment.png"),
    ("Melittin vs melittin+morphine  |  % up/down by area", FIG / "fig_mel_vs_morphine_regions.png"),
    ("Which areas change more  (FDR, recording-level)", FIG / "fig_which_regions_table.png"),
    ("Morphine minus melittin  |  by Allen area", FIG / "fig_morphine_delta_by_region.png"),
]


def add_figure_slide(prs, title, img_path):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    box = slide.shapes.add_textbox(Inches(0.35), Inches(0.12), Inches(12.6), Inches(0.38))
    p = box.text_frame.paragraphs[0]
    p.text = title
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = NAVY
    p.font.name = "Calibri"
    p.alignment = PP_ALIGN.LEFT

    img_path = Path(img_path)
    if not img_path.is_file():
        miss = slide.shapes.add_textbox(Inches(0.5), Inches(3.2), Inches(12.3), Inches(0.4))
        mp = miss.text_frame.paragraphs[0]
        mp.text = f"Missing: {img_path.name}"
        mp.font.size = Pt(14)
        mp.font.color.rgb = GREY
        return

    with Image.open(img_path) as im:
        w, h = im.size
    ar = w / float(h)
    max_w, max_h = 12.6, 6.75
    width, height = max_w, max_w / ar
    if height > max_h:
        height = max_h
        width = height * ar
    slide.shapes.add_picture(
        str(img_path),
        Inches(0.35 + (max_w - width) / 2.0),
        Inches(0.55 + (max_h - height) / 2.0),
        Inches(width),
        Inches(height),
    )


def main():
    n_before = len(Presentation(str(SRC)).slides)
    prs = Presentation(str(SRC))
    for title, path in FIGURES:
        add_figure_slide(prs, title, path)
    prs.save(str(SRC))
    extra = Path(str(SRC).replace("Mark.pptx", "Mark_which.pptx"))
    if extra.is_file() and extra.resolve() != SRC.resolve():
        extra.unlink()
    print(f"Appended {len(FIGURES)} figure slides to {SRC}")
    print(f"slides {n_before} -> {n_before + len(FIGURES)}")


if __name__ == "__main__":
    main()
